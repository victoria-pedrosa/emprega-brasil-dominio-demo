import pdfplumber
import pandas as pd
import re
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv
load_dotenv()  # lê o .env local (não vai para o GitHub)

# =========================================================================
# CONFIGURAÇÕES DE CAMINHOS DOS ARQUIVOS
# =========================================================================
CAMINHO_PASTA = os.getenv("CAMINHO_PASTA")
ARQUIVO_PDF = os.path.join(CAMINHO_PASTA, "Empréstimos Consignados - Por mês.pdf")
ARQUIVO_EXCEL = os.path.join(CAMINHO_PASTA, "CONSULTA EMPREGA.xlsx")
ARQUIVO_CSV = os.path.join(CAMINHO_PASTA, "AppSheet.ViewData.2026-08-20.csv")

# =========================================================================
# FUNÇÕES DE EXTRAÇÃO E CRUZAMENTO
# =========================================================================
def extrair_dados_pdf(caminho_pdf):
    texto_completo = ""
    with pdfplumber.open(caminho_pdf) as pdf:
        for page in pdf.pages:
            texto = page.extract_text()
            if texto:
                texto_completo += texto + "\n"
    
    blocos = texto_completo.split('Empresa:')
    dados = []
    
    for bloco in blocos[1:]:
        match_cnpj = re.search(r'([0-9]{2}\.[0-9]{3}\.[0-9]{3}/[0-9]{4}-[0-9]{2})', bloco)
        if not match_cnpj: continue
        cnpj = match_cnpj.group(1).strip()
        
        # Pega o nome da empresa e limpa a palavra 'Página: X/Y'
        nome = bloco.split('C.N.P.J.')[0]
        nome = re.sub(r'Página:.*', '', nome, flags=re.IGNORECASE)
        nome = nome.replace('|', '').replace('\n', ' ').strip()
        nome = re.sub(r'\s+', ' ', nome)
        
        dados.append({'Empresa': nome, 'CNPJ': cnpj, 'Situação': 'Com Empréstimo'})
        
    return pd.DataFrame(dados).drop_duplicates(subset=['CNPJ'])

def carregar_base_responsaveis(caminho_csv):
    # Carrega a Base de Responsáveis do CSV
    df_csv = pd.read_csv(caminho_csv, sep=';', encoding='utf-8')
    df_csv.columns = [c.replace('"', '').strip() for c in df_csv.columns]
    
    df_responsaveis = df_csv[['CNPJ:', 'Folha Responsável:']].copy()
    df_responsaveis.rename(columns={'CNPJ:': 'CNPJ', 'Folha Responsável:': 'Responsável'}, inplace=True)
    df_responsaveis['CNPJ'] = df_responsaveis['CNPJ'].astype(str).str.strip()
    
    return df_responsaveis

def gerar_planilha_formatada(caminho_excel, df_final, nome_aba):
    wb = load_workbook(caminho_excel)
    
    if nome_aba in wb.sheetnames:
        del wb[nome_aba]
    ws = wb.create_sheet(title=nome_aba)
    
    # --- 1. CRIAÇÃO DO PAINEL DE ACOMPANHAMENTO ---
    ws.merge_cells("C1:E1")
    ws["C1"] = "ACOMPANHAMENTO"
    ws["C1"].font = Font(color="008000", bold=True, size=12)
    ws["C1"].alignment = Alignment(horizontal="center")
    
    border_bottom = Border(bottom=Side(style='medium', color="00B050"))
    ws["C1"].border, ws["D1"].border, ws["E1"].border = border_bottom, border_bottom, border_bottom

    ws["C2"] = "Sem consulta"
    ws["D2"] = "Com crédito"
    ws["E2"] = "Sem Procuração"
    
    qtd_com_credito = len(df_final)
    linha_fim = 5 + qtd_com_credito # Linha onde a tabela termina
    
    ws["C3"] = 0
    ws["D3"] = qtd_com_credito
    
    # Injeta a fórmula do Excel para contar tudo que for diferente de "SIM"
    # A fórmula no Excel será: =D3 - CONT.SE(E6:Ex, "SIM")
    ws["E3"] = f'=D3-COUNTIF(E6:E{linha_fim}, "SIM")'
    
    fill_cinza = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    font_verde = Font(color="00B050", bold=True)
    
    for col in ["C", "D", "E"]:
        ws[f"{col}2"].fill = fill_cinza
        ws[f"{col}2"].alignment = Alignment(horizontal="center")
        ws[f"{col}3"].fill = fill_cinza
        ws[f"{col}3"].font = font_verde
        ws[f"{col}3"].alignment = Alignment(horizontal="center")

    # --- 2. CRIAÇÃO DA TABELA DE DADOS ---
    linha_inicio_tabela = 5
    headers = ["Empresa", "CNPJ", "Responsável", "Situação", "Procuração"]
    
    fill_cabecalho = PatternFill(start_color="364152", end_color="364152", fill_type="solid")
    font_branca = Font(color="FFFFFF", bold=True)
    borda_toda = Border(left=Side(style="thin", color="D0D7DE"), right=Side(style="thin", color="D0D7DE"),
                        top=Side(style="thin", color="D0D7DE"), bottom=Side(style="thin", color="D0D7DE"))

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=linha_inicio_tabela, column=col_num, value=header)
        cell.fill = fill_cabecalho
        cell.font = font_branca
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = borda_toda

    for i, row in df_final.iterrows():
        row_num = linha_inicio_tabela + 1 + i
        # Adiciona os dados, mas deixa a Procuração em branco ("")
        linha_dados = [row['Empresa'], row['CNPJ'], row['Responsável'], row['Situação'], ""]
        
        for col_num, valor in enumerate(linha_dados, 1):
            cell = ws.cell(row=row_num, column=col_num, value=valor)
            cell.border = borda_toda
            cell.alignment = Alignment(vertical="center", horizontal="left" if col_num == 1 else "center")

    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 15

    wb.save(caminho_excel)
    print(f"Sucesso! Aba '{nome_aba}' atualizada no arquivo {caminho_excel}.")

# =========================================================================
# FLUXO PRINCIPAL DE EXECUÇÃO
# =========================================================================
if __name__ == "__main__":
    df_pdf = extrair_dados_pdf(ARQUIVO_PDF)
    
    if not df_pdf.empty:
        # Puxa apenas a base do CSV para a coluna "Responsável"
        df_resp = carregar_base_responsaveis(ARQUIVO_CSV)
        
        # Faz o cruzamento das empresas
        df_final = pd.merge(df_pdf, df_resp, on='CNPJ', how='left')
        df_final['Responsável'] = df_final['Responsável'].fillna('Sem Responsável')
        
        # Mude o nome da aba para o formato que você preferir
        nome_aba_mes = "08.2026"
        gerar_planilha_formatada(ARQUIVO_EXCEL, df_final, nome_aba_mes)
    else:
        print("Nenhum CNPJ encontrado no arquivo PDF.")
